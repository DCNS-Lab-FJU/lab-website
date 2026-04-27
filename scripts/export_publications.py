#!/usr/bin/env python3
"""
export_publications.py
Exports data/publications.yaml → exports/publications.csv (and .xlsx if openpyxl is available)
Run from repo root: python scripts/export_publications.py
"""
import csv
import sys
from pathlib import Path

try:
    import yaml
except ImportError:
    sys.exit("Please install PyYAML: pip install pyyaml")

YAML_PATH   = Path("data/publications.yaml")
OUT_DIR     = Path("exports")
CSV_PATH    = OUT_DIR / "publications.csv"
XLSX_PATH   = OUT_DIR / "publications.xlsx"

FIELDNAMES = ["type", "year", "authors", "title", "venue", "link"]


def load_entries(data: dict) -> list[dict]:
    rows = []
    for entry in (data.get("journal_articles") or []):
        rows.append({"type": "Journal Article", **entry})
    for entry in (data.get("conference_papers") or []):
        rows.append({"type": "Conference Paper", **entry})
    for entry in (data.get("patents") or []):
        rows.append({"type": "Patent", **entry})
    for entry in (data.get("technical_reports") or []):
        rows.append({"type": "Technical Report", **entry})
    return rows


def write_csv(rows: list[dict]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with CSV_PATH.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in FIELDNAMES})
    print(f"CSV  → {CSV_PATH}  ({len(rows)} rows)")


def write_xlsx(rows: list[dict]) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("Skipping .xlsx (openpyxl not installed). Run: pip install openpyxl")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publications"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    ws.append(FIELDNAMES)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(f, "") for f in FIELDNAMES])

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    print(f"XLSX → {XLSX_PATH}  ({len(rows)} rows)")


def main():
    if not YAML_PATH.exists():
        sys.exit(f"Cannot find {YAML_PATH}. Run from repo root.")

    with YAML_PATH.open(encoding="utf-8") as f:
        data = yaml.safe_load(f)

    rows = load_entries(data)
    write_csv(rows)
    write_xlsx(rows)


if __name__ == "__main__":
    main()
